"""
API del Facturador y versión web (static/index.html).

Módulos: empresa, clientes (con estado de cuenta), proveedores (con saldo por
pagar), productos (inventario y categorías), facturas (descuento por línea,
estado de pago, anulación y devoluciones parciales),
compras a proveedores, gastos, cierre de caja, usuarios, registro de
actividad, resumen de ventas y copias de seguridad.

Acceso:
- Cada empleado entra con su usuario y contraseña (POST /api/login) y recibe
  un token que envía en el header "Authorization: Bearer <token>".
- La clave maestra (FACTURADOR_API_KEY, header "X-API-Key") da acceso de
  administrador. Es solo para el dueño: crear los primeros usuarios o
  recuperar el acceso.

Despliegue gratuito sugerido: PythonAnywhere (ver DEPLOY.md).
"""
import base64
import csv
import datetime
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import sqlite3
import sys
import tempfile
import time
import unicodedata
import uuid
from contextlib import contextmanager
from functools import wraps
from pathlib import Path

from flask import Flask, g, request, jsonify, send_file, send_from_directory
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash, generate_password_hash

SERVER_DIR = Path(__file__).parent
if str(SERVER_DIR) not in sys.path:
    sys.path.insert(0, str(SERVER_DIR))
from pdf_gen import generate_invoice_pdf, generate_statement_pdf  # noqa: E402

# -------- configuración --------
API_KEY = os.environ.get("FACTURADOR_API_KEY", "cambia-esta-clave")
DB_PATH = os.environ.get("FACTURADOR_DB", str(SERVER_DIR / "facturador_server.db"))
BACKUP_DIR = Path(os.environ.get("FACTURADOR_BACKUPS", str(Path(DB_PATH).parent / "backups")))
STATIC_DIR = SERVER_DIR / "static"

SESSION_DAYS = 30            # una sesión sin uso durante este tiempo expira
AUTO_BACKUP_EVERY = 24 * 3600
AUTO_BACKUP_KEEP = 14        # copias automáticas que se conservan
MAX_FAILED_LOGINS = 5
LOCK_SECONDS = 5 * 60
PAYMENT_METHODS = ("efectivo", "transferencia", "tarjeta", "otro")
ROLES = ("admin", "vendedor")

app = Flask(__name__, static_folder=None)

SCHEMA = """
CREATE TABLE IF NOT EXISTS company (
    id INTEGER PRIMARY KEY CHECK (id = 1),
    name TEXT DEFAULT '', tax_id TEXT DEFAULT '', address TEXT DEFAULT '',
    phone TEXT DEFAULT '', email TEXT DEFAULT '', website TEXT DEFAULT '',
    currency TEXT DEFAULT '$', tax_label TEXT DEFAULT 'IVA', tax_rate REAL DEFAULT 19,
    footer_note TEXT DEFAULT 'Gracias por su compra.',
    logo_blob BLOB, logo_mime TEXT, invoice_counter INTEGER DEFAULT 1,
    purchase_counter INTEGER DEFAULT 1, phone_country TEXT DEFAULT '57'
);
CREATE TABLE IF NOT EXISTS clients (
    id TEXT PRIMARY KEY, name TEXT NOT NULL, doc TEXT DEFAULT '', address TEXT DEFAULT '',
    phone TEXT DEFAULT '', email TEXT DEFAULT '', created_at REAL
);
CREATE TABLE IF NOT EXISTS suppliers (
    id TEXT PRIMARY KEY, name TEXT NOT NULL, contact_name TEXT DEFAULT '',
    phone TEXT DEFAULT '', email TEXT DEFAULT '', created_at REAL
);
CREATE TABLE IF NOT EXISTS products (
    id TEXT PRIMARY KEY, name TEXT NOT NULL, price REAL NOT NULL,
    unit TEXT DEFAULT 'unidad', supplier_id TEXT, stock_qty REAL DEFAULT 0,
    cost_price REAL DEFAULT 0, created_at REAL, category TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS invoices (
    id TEXT PRIMARY KEY, number TEXT UNIQUE NOT NULL, created_at REAL,
    client_name TEXT, client_doc TEXT, client_address TEXT, client_phone TEXT, client_email TEXT,
    tax_rate REAL DEFAULT 0, notes TEXT DEFAULT '', warranty TEXT DEFAULT '', seller TEXT DEFAULT '',
    status TEXT DEFAULT 'activa', voided_at REAL, voided_by TEXT DEFAULT '', void_reason TEXT DEFAULT '',
    user_id TEXT, user_name TEXT DEFAULT '', payment_tracking INTEGER DEFAULT 0,
    client_id TEXT, public_token TEXT, tax_included INTEGER DEFAULT 0
);
CREATE TABLE IF NOT EXISTS invoice_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    invoice_id TEXT NOT NULL REFERENCES invoices(id) ON DELETE CASCADE,
    product_id TEXT, description TEXT NOT NULL, qty REAL NOT NULL, price REAL NOT NULL,
    discount REAL DEFAULT 0, unit_cost REAL
);
CREATE TABLE IF NOT EXISTS stock_movements (
    id INTEGER PRIMARY KEY AUTOINCREMENT, product_id TEXT NOT NULL, change_qty REAL NOT NULL,
    reason TEXT NOT NULL, reference TEXT, created_at REAL
);
CREATE TABLE IF NOT EXISTS users (
    id TEXT PRIMARY KEY, username TEXT UNIQUE NOT NULL COLLATE NOCASE, name TEXT NOT NULL,
    password_hash TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'vendedor', active INTEGER DEFAULT 1,
    failed_logins INTEGER DEFAULT 0, locked_until REAL DEFAULT 0, created_at REAL
);
CREATE TABLE IF NOT EXISTS sessions (
    token_hash TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
    created_at REAL, last_seen REAL
);
CREATE TABLE IF NOT EXISTS payments (
    id TEXT PRIMARY KEY, invoice_id TEXT NOT NULL REFERENCES invoices(id) ON DELETE CASCADE,
    amount REAL NOT NULL, method TEXT NOT NULL, note TEXT DEFAULT '', created_at REAL,
    user_id TEXT, user_name TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS returns (
    id TEXT PRIMARY KEY, invoice_id TEXT NOT NULL REFERENCES invoices(id) ON DELETE CASCADE,
    amount REAL NOT NULL, net REAL NOT NULL, cost REAL DEFAULT 0, refund REAL DEFAULT 0,
    refund_method TEXT, reason TEXT DEFAULT '', created_at REAL, user_id TEXT, user_name TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS return_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT, return_id TEXT NOT NULL REFERENCES returns(id) ON DELETE CASCADE,
    invoice_item_id INTEGER NOT NULL, product_id TEXT, description TEXT, qty REAL NOT NULL, amount REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS purchases (
    id TEXT PRIMARY KEY, number TEXT UNIQUE NOT NULL, supplier_id TEXT, supplier_name TEXT DEFAULT '',
    supplier_ref TEXT DEFAULT '', notes TEXT DEFAULT '', created_at REAL, user_id TEXT, user_name TEXT DEFAULT '',
    status TEXT DEFAULT 'activa', voided_at REAL, voided_by TEXT DEFAULT '', void_reason TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS purchase_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT, purchase_id TEXT NOT NULL REFERENCES purchases(id) ON DELETE CASCADE,
    product_id TEXT, description TEXT NOT NULL, qty REAL NOT NULL, unit_cost REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS supplier_payments (
    id TEXT PRIMARY KEY, purchase_id TEXT NOT NULL REFERENCES purchases(id) ON DELETE CASCADE,
    amount REAL NOT NULL, method TEXT NOT NULL, note TEXT DEFAULT '', created_at REAL,
    user_id TEXT, user_name TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS expenses (
    id TEXT PRIMARY KEY, spent_at REAL NOT NULL, category TEXT DEFAULT '', description TEXT DEFAULT '',
    amount REAL NOT NULL, method TEXT NOT NULL, supplier_id TEXT, created_at REAL,
    user_id TEXT, user_name TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS cash_closings (
    id TEXT PRIMARY KEY, day TEXT NOT NULL, scope_user_id TEXT, scope_name TEXT DEFAULT '',
    opening_cash REAL DEFAULT 0, expected TEXT NOT NULL, counted TEXT NOT NULL, difference REAL DEFAULT 0,
    notes TEXT DEFAULT '', created_at REAL, user_id TEXT, user_name TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS activity_log (
    id INTEGER PRIMARY KEY AUTOINCREMENT, created_at REAL NOT NULL, user_id TEXT, user_name TEXT DEFAULT '',
    entity TEXT DEFAULT '', entity_id TEXT, summary TEXT NOT NULL
);
"""

# columnas que se agregaron después del primer lanzamiento: se aplican con
# ALTER TABLE sobre bases de datos ya existentes, sin perder datos.
MIGRATIONS = [
    ("products", "supplier_id", "ALTER TABLE products ADD COLUMN supplier_id TEXT"),
    ("products", "stock_qty", "ALTER TABLE products ADD COLUMN stock_qty REAL DEFAULT 0"),
    ("products", "cost_price", "ALTER TABLE products ADD COLUMN cost_price REAL DEFAULT 0"),
    ("products", "category", "ALTER TABLE products ADD COLUMN category TEXT DEFAULT ''"),
    ("company", "purchase_counter", "ALTER TABLE company ADD COLUMN purchase_counter INTEGER DEFAULT 1"),
    ("company", "phone_country", "ALTER TABLE company ADD COLUMN phone_country TEXT DEFAULT '57'"),
    ("invoices", "warranty", "ALTER TABLE invoices ADD COLUMN warranty TEXT DEFAULT ''"),
    ("invoices", "seller", "ALTER TABLE invoices ADD COLUMN seller TEXT DEFAULT ''"),
    ("invoices", "status", "ALTER TABLE invoices ADD COLUMN status TEXT DEFAULT 'activa'"),
    ("invoices", "voided_at", "ALTER TABLE invoices ADD COLUMN voided_at REAL"),
    ("invoices", "voided_by", "ALTER TABLE invoices ADD COLUMN voided_by TEXT DEFAULT ''"),
    ("invoices", "void_reason", "ALTER TABLE invoices ADD COLUMN void_reason TEXT DEFAULT ''"),
    ("invoices", "user_id", "ALTER TABLE invoices ADD COLUMN user_id TEXT"),
    ("invoices", "user_name", "ALTER TABLE invoices ADD COLUMN user_name TEXT DEFAULT ''"),
    # Las facturas anteriores al registro de pagos quedan con 0 ("sin registro de pago"),
    # para no marcarlas como pendientes ni como pagadas sin saberlo.
    ("invoices", "payment_tracking", "ALTER TABLE invoices ADD COLUMN payment_tracking INTEGER DEFAULT 0"),
    ("invoices", "client_id", "ALTER TABLE invoices ADD COLUMN client_id TEXT"),
    # public_token: enlaces públicos de facturas, ya retirados (la columna queda sin uso)
    ("invoices", "public_token", "ALTER TABLE invoices ADD COLUMN public_token TEXT"),
    # 1 = el precio de los artículos ya trae el IVA (se desglosa); 0 = facturas viejas, el IVA se sumaba aparte.
    ("invoices", "tax_included", "ALTER TABLE invoices ADD COLUMN tax_included INTEGER DEFAULT 0"),
    ("invoice_items", "product_id", "ALTER TABLE invoice_items ADD COLUMN product_id TEXT"),
    ("invoice_items", "discount", "ALTER TABLE invoice_items ADD COLUMN discount REAL DEFAULT 0"),
    ("invoice_items", "unit_cost", "ALTER TABLE invoice_items ADD COLUMN unit_cost REAL"),
]

INDEXES = """
CREATE INDEX IF NOT EXISTS idx_items_invoice ON invoice_items(invoice_id);
CREATE INDEX IF NOT EXISTS idx_payments_invoice ON payments(invoice_id);
CREATE INDEX IF NOT EXISTS idx_payments_created ON payments(created_at);
CREATE INDEX IF NOT EXISTS idx_invoices_created ON invoices(created_at);
CREATE INDEX IF NOT EXISTS idx_invoices_client ON invoices(client_id);
CREATE UNIQUE INDEX IF NOT EXISTS idx_invoices_token ON invoices(public_token);
CREATE INDEX IF NOT EXISTS idx_sessions_user ON sessions(user_id);
CREATE INDEX IF NOT EXISTS idx_returns_invoice ON returns(invoice_id);
CREATE INDEX IF NOT EXISTS idx_purchases_supplier ON purchases(supplier_id);
CREATE INDEX IF NOT EXISTS idx_purchase_items ON purchase_items(purchase_id);
CREATE INDEX IF NOT EXISTS idx_supplier_payments ON supplier_payments(purchase_id);
CREATE INDEX IF NOT EXISTS idx_expenses_spent ON expenses(spent_at);
CREATE UNIQUE INDEX IF NOT EXISTS idx_closings_scope ON cash_closings(day, COALESCE(scope_user_id, ''));
CREATE INDEX IF NOT EXISTS idx_activity_created ON activity_log(created_at);
"""


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(SCHEMA)
    conn.execute("INSERT OR IGNORE INTO company (id) VALUES (1)")
    for table, col, ddl in MIGRATIONS:
        cols = {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
        if col not in cols:
            conn.execute(ddl)
    conn.executescript(INDEXES)
    # Las facturas viejas guardaban solo el nombre del cliente: se enlazan con su
    # ficha (mismo nombre y documento) para que aparezcan en su estado de cuenta.
    conn.execute(
        """UPDATE invoices SET client_id = (
             SELECT c.id FROM clients c WHERE c.name = invoices.client_name
               AND COALESCE(c.doc, '') = COALESCE(invoices.client_doc, '') LIMIT 1)
           WHERE client_id IS NULL""")
    _fix_voided_without_refund(conn)
    conn.commit()
    conn.close()


def _fix_voided_without_refund(conn):
    """Antes, anular una factura (o compra) pagada no registraba la devolución del dinero y la caja lo
    seguía contando. Se registra ahora, con la fecha de la anulación y por el mismo método con que se pagó.
    Es idempotente: tras el ajuste el neto pagado queda en cero y no se vuelve a aplicar."""
    fixes = [("invoices", "payments", "invoice_id", "factura", "Dinero devuelto al anular la factura (ajuste automático)"),
             ("purchases", "supplier_payments", "purchase_id", "compra",
              "Dinero devuelto por el proveedor al anular la compra (ajuste automático)")]
    for table, pay_table, key, entity, note in fixes:
        for r in conn.execute(f"SELECT id, number, voided_at, created_at FROM {table} WHERE status = 'anulada'").fetchall():
            split = _refund_split(conn, pay_table, key, r[0])
            if not split:
                continue
            for method, amount in split.items():
                conn.execute(
                    f"""INSERT INTO {pay_table} (id, {key}, amount, method, note, created_at, user_id, user_name)
                        VALUES (?,?,?,?,?,?,NULL,'Ajuste automático')""",
                    (str(uuid.uuid4()), r[0], -amount, method, note, r[2] or r[3] or time.time()))
            conn.execute(
                "INSERT INTO activity_log (created_at, user_id, user_name, entity, entity_id, summary) VALUES (?,NULL,?,?,?,?)",
                (time.time(), "Ajuste automático", entity, r[0],
                 f"Registró la devolución de {_split_text(split)} de la {entity} anulada {r[1]}, "
                 "que antes seguía contando en la caja"))



@contextmanager
def get_conn():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.execute("PRAGMA busy_timeout=8000")
    conn.execute("PRAGMA foreign_keys=ON")
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


class ApiError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.message = message
        self.status = status


@app.errorhandler(ApiError)
def handle_api_error(e):
    return jsonify({"error": e.message}), e.status


@app.errorhandler(Exception)
def handle_any_error(e):
    if isinstance(e, HTTPException):
        return e
    app.logger.exception("Error no controlado")
    return jsonify({"error": f"Error en el servidor: {e}"}), 500


def money(x):
    return round((x or 0) + 1e-9, 2)


def fmt(x):
    return f"{money(x):,.2f}"


def num(value, field, minimum=None, maximum=None, default=None):
    """Convierte a número validando rango; lanza ApiError con mensaje claro."""
    if value is None or value == "":
        if default is not None:
            return default
        raise ApiError(f"Falta el valor de {field}.")
    try:
        v = float(value)
    except (TypeError, ValueError):
        raise ApiError(f"{field} debe ser un número.")
    if v != v or v in (float("inf"), float("-inf")):
        raise ApiError(f"{field} debe ser un número.")
    if minimum is not None and v < minimum:
        raise ApiError(f"{field} no puede ser menor que {minimum:g}.")
    if maximum is not None and v > maximum:
        raise ApiError(f"{field} no puede ser mayor que {maximum:g}.")
    return v


def body():
    d = request.get_json(force=True, silent=True)
    if not isinstance(d, dict):
        raise ApiError("La petición no trae datos válidos.")
    return d


def text(d, key, default=""):
    v = d.get(key, default)
    return default if v is None else str(v).strip()


def parse_method(method):
    method = (method or "efectivo").strip().lower()
    if method not in PAYMENT_METHODS:
        raise ApiError("Método de pago no válido. Usa: " + ", ".join(PAYMENT_METHODS) + ".")
    return method


def tz_offset_arg():
    return num(request.args.get("tz"), "tz", -14 * 60, 14 * 60, default=0)  # minutos, como JS getTimezoneOffset()


def day_bounds(day, tz_offset):
    """Inicio y fin (epoch) de un día local 'YYYY-MM-DD' dado el desfase del navegador."""
    try:
        d = datetime.date.fromisoformat(day)
    except (TypeError, ValueError):
        raise ApiError("La fecha no es válida (usa AAAA-MM-DD).")
    start = datetime.datetime(d.year, d.month, d.day, tzinfo=datetime.timezone.utc).timestamp() + tz_offset * 60
    return start, start + 86400


# ---------- registro de actividad ----------
def log(conn, entity, entity_id, summary):
    user = getattr(g, "user", None) or {}
    conn.execute(
        "INSERT INTO activity_log (created_at, user_id, user_name, entity, entity_id, summary) VALUES (?,?,?,?,?,?)",
        (time.time(), user.get("id"), user.get("name", ""), entity, entity_id, summary))


# ---------- autenticación ----------
MASTER_USER = {"id": None, "username": "clave-maestra", "name": "Administrador", "role": "admin", "master": True}
_DUMMY_HASH = generate_password_hash("contraseña-de-relleno")


def _token_hash(token):
    return hashlib.sha256(token.encode()).hexdigest()


def _current_user():
    key = request.headers.get("X-API-Key", "")
    if key:
        if hmac.compare_digest(key.encode(), API_KEY.encode()):
            return MASTER_USER
        return None
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None
    th = _token_hash(auth[7:].strip())
    now = time.time()
    with get_conn() as conn:
        row = conn.execute(
            """SELECT u.id, u.username, u.name, u.role, u.active, s.last_seen FROM sessions s
               JOIN users u ON u.id = s.user_id WHERE s.token_hash=?""", (th,)).fetchone()
        if not row or not row["active"] or (row["last_seen"] or 0) < now - SESSION_DAYS * 86400:
            return None
        if now - (row["last_seen"] or 0) > 3600:
            conn.execute("UPDATE sessions SET last_seen=? WHERE token_hash=?", (now, th))
    return {"id": row["id"], "username": row["username"], "name": row["name"], "role": row["role"], "master": False}


def require_auth(admin=False):
    def deco(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            user = _current_user()
            if not user:
                return jsonify({"error": "Sesión vencida o datos de acceso incorrectos. Vuelve a entrar."}), 401
            if admin and user["role"] != "admin":
                return jsonify({"error": "Solo un administrador puede hacer esto."}), 403
            g.user = user
            return fn(*args, **kwargs)
        return wrapper
    return deco


def is_admin():
    return g.user["role"] == "admin"


def public_user(row):
    return {k: row[k] for k in ("id", "username", "name", "role")} | {"active": bool(row["active"])}


@app.post("/api/login")
def login():
    d = body()
    username = text(d, "username")
    password = str(d.get("password", ""))
    if not username or not password:
        raise ApiError("Escribe tu usuario y tu contraseña.")
    now = time.time()
    with get_conn() as conn:
        u = conn.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        if not u:
            check_password_hash(_DUMMY_HASH, password)  # mismo tiempo de respuesta
            raise ApiError("Usuario o contraseña incorrectos.", 401)
        g.user = {"id": u["id"], "name": u["name"]}
        if (u["locked_until"] or 0) > now:
            mins = int((u["locked_until"] - now) // 60) + 1
            raise ApiError(f"Demasiados intentos fallidos. Intenta de nuevo en {mins} minuto(s).", 429)
        if not check_password_hash(u["password_hash"], password):
            fails = (u["failed_logins"] or 0) + 1
            if fails >= MAX_FAILED_LOGINS:
                conn.execute("UPDATE users SET failed_logins=0, locked_until=? WHERE id=?", (now + LOCK_SECONDS, u["id"]))
                log(conn, "usuario", u["id"], f"Usuario bloqueado 5 minutos por {MAX_FAILED_LOGINS} intentos fallidos")
            else:
                conn.execute("UPDATE users SET failed_logins=? WHERE id=?", (fails, u["id"]))
            conn.commit()
            raise ApiError("Usuario o contraseña incorrectos.", 401)
        if not u["active"]:
            raise ApiError("Este usuario está desactivado. Habla con el administrador.", 403)
        token = secrets.token_urlsafe(32)
        conn.execute("UPDATE users SET failed_logins=0, locked_until=0 WHERE id=?", (u["id"],))
        conn.execute("DELETE FROM sessions WHERE last_seen < ?", (now - SESSION_DAYS * 86400,))
        conn.execute("INSERT INTO sessions (token_hash, user_id, created_at, last_seen) VALUES (?,?,?,?)",
                     (_token_hash(token), u["id"], now, now))
        log(conn, "usuario", u["id"], "Entró al sistema")
        return jsonify({"token": token, "user": public_user(u)})


@app.post("/api/logout")
@require_auth()
def logout():
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        with get_conn() as conn:
            conn.execute("DELETE FROM sessions WHERE token_hash=?", (_token_hash(auth[7:].strip()),))
    return jsonify({"ok": True})


@app.get("/api/me")
@require_auth()
def me():
    return jsonify(dict(g.user))


@app.post("/api/me/password")
@require_auth()
def change_own_password():
    if g.user["master"]:
        raise ApiError("La clave maestra se cambia en el servidor, no desde aquí.")
    d = body()
    new = str(d.get("new_password", ""))
    _check_password_rules(new)
    with get_conn() as conn:
        u = conn.execute("SELECT password_hash FROM users WHERE id=?", (g.user["id"],)).fetchone()
        if not check_password_hash(u["password_hash"], str(d.get("current_password", ""))):
            raise ApiError("La contraseña actual no es correcta.")
        conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new), g.user["id"]))
        # cierra las demás sesiones de este usuario (p. ej. un celular perdido)
        auth = request.headers.get("Authorization", "")
        conn.execute("DELETE FROM sessions WHERE user_id=? AND token_hash<>?",
                     (g.user["id"], _token_hash(auth[7:].strip())))
        log(conn, "usuario", g.user["id"], "Cambió su contraseña")
    return jsonify({"ok": True})


def _check_password_rules(pw):
    if len(pw) < 6:
        raise ApiError("La contraseña debe tener al menos 6 caracteres.")


# ---------- usuarios (solo administrador) ----------
@app.get("/api/users")
@require_auth(admin=True)
def list_users():
    with get_conn() as conn:
        rows = conn.execute("SELECT * FROM users ORDER BY active DESC, name COLLATE NOCASE").fetchall()
        return jsonify([public_user(r) for r in rows])


@app.post("/api/users")
@require_auth(admin=True)
def create_user():
    d = body()
    username, name, role = text(d, "username"), text(d, "name"), d.get("role", "vendedor")
    password = str(d.get("password", ""))
    if not username or not name:
        raise ApiError("El usuario necesita nombre y nombre de usuario.")
    if " " in username:
        raise ApiError("El nombre de usuario no puede tener espacios.")
    if role not in ROLES:
        raise ApiError("El rol debe ser 'admin' o 'vendedor'.")
    _check_password_rules(password)
    uid = str(uuid.uuid4())
    try:
        with get_conn() as conn:
            conn.execute(
                "INSERT INTO users (id, username, name, password_hash, role, active, created_at) VALUES (?,?,?,?,?,1,?)",
                (uid, username, name, generate_password_hash(password), role, time.time()))
            log(conn, "usuario", uid, f'Creó el usuario "{username}" ({name}) como {role}')
    except sqlite3.IntegrityError:
        raise ApiError(f'Ya existe un usuario "{username}".')
    return jsonify({"id": uid})


@app.put("/api/users/<uid>")
@require_auth(admin=True)
def update_user(uid):
    d = body()
    with get_conn() as conn:
        u = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
        if not u:
            raise ApiError("Usuario no encontrado.", 404)
        name = text(d, "name", u["name"]) or u["name"]
        role = d.get("role", u["role"])
        active = 1 if d.get("active", bool(u["active"])) else 0
        if role not in ROLES:
            raise ApiError("El rol debe ser 'admin' o 'vendedor'.")
        if uid == g.user["id"] and (not active or role != "admin"):
            raise ApiError("No puedes quitarte a ti mismo el rol de administrador ni desactivarte.")
        conn.execute("UPDATE users SET name=?, role=?, active=? WHERE id=?", (name, role, active, uid))
        changes = []
        if name != u["name"]:
            changes.append(f'nombre "{u["name"]}" → "{name}"')
        if role != u["role"]:
            changes.append(f"rol {u['role']} → {role}")
        if active != u["active"]:
            changes.append("activado" if active else "desactivado")
        if d.get("password"):
            _check_password_rules(str(d["password"]))
            conn.execute("UPDATE users SET password_hash=?, failed_logins=0, locked_until=0 WHERE id=?",
                         (generate_password_hash(str(d["password"])), uid))
            changes.append("contraseña cambiada")
        if not active or d.get("password"):
            conn.execute("DELETE FROM sessions WHERE user_id=?", (uid,))
        if changes:
            log(conn, "usuario", uid, f'Editó el usuario "{u["username"]}": ' + ", ".join(changes))
    return jsonify({"ok": True})


# ---------- versión web (celular / navegador) ----------
@app.get("/")
def web_index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.get("/manifest.webmanifest")
def web_manifest():
    return send_from_directory(STATIC_DIR, "manifest.webmanifest", mimetype="application/manifest+json")


@app.get("/icon.svg")
def web_icon():
    return send_from_directory(STATIC_DIR, "icon.svg", mimetype="image/svg+xml")


# ---------- salud ----------
@app.get("/api/ping")
@require_auth()
def ping():
    return jsonify({"ok": True, "name": "Facturador API"})


# ---------- empresa ----------
def _company(conn):
    return dict(conn.execute("SELECT * FROM company WHERE id=1").fetchone())


@app.get("/api/company")
@require_auth()
def get_company():
    with get_conn() as conn:
        data = _company(conn)
    logo_blob = data.pop("logo_blob", None)
    data["logo_base64"] = base64.b64encode(logo_blob).decode("ascii") if logo_blob else None
    return jsonify(data)


@app.post("/api/company")
@require_auth(admin=True)
def save_company():
    d = body()
    tax_rate = num(d.get("tax_rate"), "La tarifa de impuesto", 0, 100, default=0)
    country = "".join(ch for ch in text(d, "phone_country", "57") if ch.isdigit()) or "57"
    with get_conn() as conn:
        conn.execute(
            """UPDATE company SET name=?, tax_id=?, address=?, phone=?, email=?, website=?,
               currency=?, tax_label=?, tax_rate=?, footer_note=?, phone_country=? WHERE id=1""",
            (text(d, "name"), text(d, "tax_id"), text(d, "address"), text(d, "phone"), text(d, "email"),
             text(d, "website"), text(d, "currency", "$") or "$", text(d, "tax_label", "IVA") or "IVA",
             tax_rate, text(d, "footer_note"), country))
        log(conn, "empresa", "1", "Editó los datos de la empresa")
    return jsonify({"ok": True})


@app.post("/api/company/logo")
@require_auth(admin=True)
def save_logo():
    d = body()
    blob = base64.b64decode(d["data_base64"]) if d.get("data_base64") else None
    if blob and len(blob) > 3 * 1024 * 1024:
        raise ApiError("El logo pesa demasiado (máximo 3 MB).")
    with get_conn() as conn:
        conn.execute("UPDATE company SET logo_blob=?, logo_mime=? WHERE id=1", (blob, d.get("mime")))
        log(conn, "empresa", "1", "Cambió el logo" if blob else "Quitó el logo")
    return jsonify({"ok": True})


# ---------- clientes ----------
@app.get("/api/clients")
@require_auth()
def list_clients():
    with get_conn() as conn:
        rows = [dict(r) for r in conn.execute("SELECT * FROM clients ORDER BY name COLLATE NOCASE").fetchall()]
        owed = {}
        for inv in _invoice_rows(conn, "WHERE i.client_id IS NOT NULL"):
            if inv["payment_status"] in ("pendiente", "abono"):
                owed[inv["client_id"]] = owed.get(inv["client_id"], 0) + inv["balance"]
    for r in rows:
        r["balance"] = money(owed.get(r["id"], 0))
    return jsonify(rows)


def _client_values(d):
    name = text(d, "name")
    if not name:
        raise ApiError("El nombre del cliente es obligatorio.")
    return (name, text(d, "doc"), text(d, "address"), text(d, "phone"), text(d, "email"))


@app.post("/api/clients")
@require_auth()
def create_client():
    vals = _client_values(body())
    cid = str(uuid.uuid4())
    with get_conn() as conn:
        conn.execute("INSERT INTO clients (name, doc, address, phone, email, id, created_at) VALUES (?,?,?,?,?,?,?)",
                     vals + (cid, time.time()))
        log(conn, "cliente", cid, f'Creó el cliente "{vals[0]}"')
    return jsonify({"id": cid})


@app.put("/api/clients/<cid>")
@require_auth()
def update_client(cid):
    vals = _client_values(body())
    with get_conn() as conn:
        old = conn.execute("SELECT * FROM clients WHERE id=?", (cid,)).fetchone()
        if not old:
            raise ApiError("Cliente no encontrado.", 404)
        conn.execute("UPDATE clients SET name=?, doc=?, address=?, phone=?, email=? WHERE id=?", vals + (cid,))
        changes = [f'{label}: "{old[k] or ""}" → "{v}"' for k, label, v in
                   zip(("name", "doc", "address", "phone", "email"),
                       ("nombre", "documento", "dirección", "teléfono", "correo"), vals) if (old[k] or "") != v]
        if changes:
            log(conn, "cliente", cid, f'Editó el cliente "{vals[0]}": ' + "; ".join(changes))
    return jsonify({"id": cid})


@app.delete("/api/clients/<cid>")
@require_auth(admin=True)
def delete_client(cid):
    with get_conn() as conn:
        old = conn.execute("SELECT name FROM clients WHERE id=?", (cid,)).fetchone()
        conn.execute("DELETE FROM clients WHERE id=?", (cid,))
        if old:
            log(conn, "cliente", cid, f'Eliminó el cliente "{old["name"]}"')
    return jsonify({"ok": True})


def _statement(conn, cid):
    client = conn.execute("SELECT * FROM clients WHERE id=?", (cid,)).fetchone()
    if not client:
        raise ApiError("Cliente no encontrado.", 404)
    invoices = _invoice_rows(conn, "WHERE i.client_id=?", (cid,), order="ORDER BY i.created_at")
    active = [i for i in invoices if i["status"] != "anulada"]
    tracked = [i for i in active if i["payment_tracking"]]
    pending = [i for i in active if i["payment_status"] in ("pendiente", "abono")]
    return {
        "client": dict(client),
        "invoices": invoices,
        "summary": {
            "count": len(active),
            "total": money(sum(i["total"] for i in active)),
            "paid": money(sum(i["paid"] for i in tracked)),
            "balance": money(sum(i["balance"] for i in pending)),
            "pending": [{"id": i["id"], "number": i["number"], "created_at": i["created_at"],
                         "balance": i["balance"]} for i in pending],
        },
    }


@app.get("/api/clients/<cid>/statement")
@require_auth()
def client_statement(cid):
    with get_conn() as conn:
        return jsonify(_statement(conn, cid))


@app.get("/api/clients/<cid>/statement.pdf")
@require_auth()
def client_statement_pdf(cid):
    with get_conn() as conn:
        st = _statement(conn, cid)
        company = _company(conn)
    buf = io.BytesIO()
    generate_statement_pdf(st, company, buf)
    buf.seek(0)
    safe = "".join(ch for ch in st["client"]["name"] if ch.isalnum() or ch in " -_")[:40].strip() or "cliente"
    return send_file(buf, mimetype="application/pdf", download_name=f"Estado de cuenta - {safe}.pdf")


# ---------- proveedores ----------
@app.get("/api/suppliers")
@require_auth()
def list_suppliers():
    with get_conn() as conn:
        rows = [dict(r) for r in conn.execute("SELECT * FROM suppliers ORDER BY name COLLATE NOCASE").fetchall()]
        owed = {}
        for p in _purchase_rows(conn, "WHERE p.status <> 'anulada'"):
            owed[p["supplier_id"]] = owed.get(p["supplier_id"], 0) + p["balance"]
    for r in rows:
        r["balance"] = money(owed.get(r["id"], 0))
    return jsonify(rows)


def _supplier_values(d):
    name = text(d, "name")
    if not name:
        raise ApiError("El nombre del proveedor es obligatorio.")
    return (name, text(d, "contact_name"), text(d, "phone"), text(d, "email"))


@app.post("/api/suppliers")
@require_auth(admin=True)
def create_supplier():
    vals = _supplier_values(body())
    sid = str(uuid.uuid4())
    with get_conn() as conn:
        conn.execute("INSERT INTO suppliers (name, contact_name, phone, email, id, created_at) VALUES (?,?,?,?,?,?)",
                     vals + (sid, time.time()))
        log(conn, "proveedor", sid, f'Creó el proveedor "{vals[0]}"')
    return jsonify({"id": sid})


@app.put("/api/suppliers/<sid>")
@require_auth(admin=True)
def update_supplier(sid):
    vals = _supplier_values(body())
    with get_conn() as conn:
        conn.execute("UPDATE suppliers SET name=?, contact_name=?, phone=?, email=? WHERE id=?", vals + (sid,))
        log(conn, "proveedor", sid, f'Editó el proveedor "{vals[0]}"')
    return jsonify({"id": sid})


@app.delete("/api/suppliers/<sid>")
@require_auth(admin=True)
def delete_supplier(sid):
    with get_conn() as conn:
        old = conn.execute("SELECT name FROM suppliers WHERE id=?", (sid,)).fetchone()
        conn.execute("DELETE FROM suppliers WHERE id=?", (sid,))
        if old:
            log(conn, "proveedor", sid, f'Eliminó el proveedor "{old["name"]}"')
    return jsonify({"ok": True})


# ---------- productos ----------
@app.get("/api/products")
@require_auth()
def list_products():
    with get_conn() as conn:
        rows = conn.execute(
            """SELECT p.*, s.name AS supplier_name FROM products p
               LEFT JOIN suppliers s ON s.id = p.supplier_id
               ORDER BY p.name COLLATE NOCASE"""
        ).fetchall()
    result = [dict(r) for r in rows]
    if not is_admin():
        for r in result:
            r.pop("cost_price", None)  # el costo (y la ganancia) solo lo ve el administrador
    return jsonify(result)


PRODUCT_FIELDS = ("name", "price", "unit", "supplier_id", "stock_qty", "cost_price", "category")


def _product_values(d):
    name = text(d, "name")
    if not name:
        raise ApiError("El producto necesita un nombre.")
    price = num(d.get("price"), "El precio", 0)
    if price <= 0:
        raise ApiError("El precio debe ser mayor que cero.")
    return (name, price, text(d, "unit") or "unidad", d.get("supplier_id") or None,
            num(d.get("stock_qty"), "Las existencias", 0, default=0),
            num(d.get("cost_price"), "El costo", 0, default=0), text(d, "category"))


@app.post("/api/products")
@require_auth(admin=True)
def create_product():
    vals = _product_values(body())
    pid = str(uuid.uuid4())
    with get_conn() as conn:
        conn.execute(
            """INSERT INTO products (name, price, unit, supplier_id, stock_qty, cost_price, category, id, created_at)
               VALUES (?,?,?,?,?,?,?,?,?)""", vals + (pid, time.time()))
        log(conn, "producto", pid, f'Creó el producto "{vals[0]}" (precio {fmt(vals[1])}, stock {vals[4]:g})')
    return jsonify({"id": pid})


@app.put("/api/products/<pid>")
@require_auth(admin=True)
def update_product(pid):
    vals = _product_values(body())
    with get_conn() as conn:
        old = conn.execute("SELECT * FROM products WHERE id=?", (pid,)).fetchone()
        if not old:
            raise ApiError("Producto no encontrado.", 404)
        conn.execute(
            "UPDATE products SET name=?, price=?, unit=?, supplier_id=?, stock_qty=?, cost_price=?, category=? WHERE id=?",
            vals + (pid,))
        labels = {"name": "nombre", "price": "precio", "unit": "unidad", "stock_qty": "stock",
                  "cost_price": "costo", "category": "categoría", "supplier_id": "proveedor"}
        changes = []
        for field, new in zip(PRODUCT_FIELDS, vals):
            prev = old[field]
            if isinstance(new, float):
                if abs((prev or 0) - new) > 1e-9:
                    changes.append(f"{labels[field]} {prev or 0:,.2f} → {new:,.2f}" if field != "stock_qty"
                                   else f"stock {prev or 0:g} → {new:g}")
            elif (prev or "") != (new or ""):
                changes.append(f"{labels[field]} cambiado" if field == "supplier_id"
                               else f'{labels[field]} "{prev or ""}" → "{new}"')
        if changes:
            if old["stock_qty"] != vals[4]:
                conn.execute(
                    "INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                    (pid, vals[4] - (old["stock_qty"] or 0), "ajuste manual", g.user["name"], time.time()))
            log(conn, "producto", pid, f'Editó el producto "{vals[0]}": ' + "; ".join(changes))
    return jsonify({"id": pid})


@app.delete("/api/products/<pid>")
@require_auth(admin=True)
def delete_product(pid):
    with get_conn() as conn:
        old = conn.execute("SELECT name FROM products WHERE id=?", (pid,)).fetchone()
        conn.execute("DELETE FROM products WHERE id=?", (pid,))
        if old:
            log(conn, "producto", pid, f'Eliminó el producto "{old["name"]}"')
    return jsonify({"ok": True})


@app.post("/api/products/<pid>/restock")
@require_auth(admin=True)
def restock_product(pid):
    d = body()
    qty = num(d.get("qty"), "La cantidad", 0)
    if qty <= 0:
        raise ApiError("La cantidad a ingresar debe ser mayor que cero.")
    with get_conn() as conn:
        row = conn.execute("SELECT name, stock_qty FROM products WHERE id=?", (pid,)).fetchone()
        if not row:
            raise ApiError("Producto no encontrado.", 404)
        new_qty = (row["stock_qty"] or 0) + qty
        conn.execute("UPDATE products SET stock_qty=? WHERE id=?", (new_qty, pid))
        conn.execute(
            "INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
            (pid, qty, "reabastecimiento", text(d, "note"), time.time()))
        log(conn, "producto", pid, f'Sumó {qty:g} al stock de "{row["name"]}" ({new_qty:g} en total)'
            + (f" — {text(d, 'note')}" if text(d, "note") else ""))
    return jsonify({"ok": True, "stock_qty": new_qty})


# ---------- importar productos desde Excel / CSV ----------
IMPORT_MAX_BYTES = 5 * 1024 * 1024
IMPORT_MAX_ROWS = 5000
IMPORT_FIELDS = {"name": "Nombre", "category": "Categoría", "price": "Precio de venta (IVA incluido)",
                 "cost_price": "Costo", "stock_qty": "Existencias", "unit": "Unidad"}


def _norm(v):
    v = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", v).strip()


def _header_field(h):
    """A qué campo corresponde un encabezado (tolera mayúsculas, tildes y variantes)."""
    h = _norm(h)
    if not h:
        return None
    if "costo" in h or "compra" in h:
        return "cost_price"
    if "precio" in h or "valor" in h or "pvp" in h or h == "venta":
        return "price"
    if any(w in h for w in ("existencia", "stock", "cantidad", "inventario", "unidades disponibles")):
        return "stock_qty"
    if any(w in h for w in ("categoria", "grupo", "familia", "linea")):
        return "category"
    if h in ("unidad", "unidad de medida", "medida", "um"):
        return "unit"
    if any(w in h for w in ("nombre", "producto", "articulo", "descripcion", "item", "referencia")):
        return "name"
    return None


def _parse_amount(v):
    """Números como los escribe la gente en Colombia: 25000, 25.000, $ 25.000, 1.234.567,50, 25,5…"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = re.sub(r"[^0-9,.\-]", "", str(v))
    if not t or t in "-.,":
        return None
    if "," in t and "." in t:
        dec = "," if t.rfind(",") > t.rfind(".") else "."
        t = t.replace("." if dec == "," else ",", "").replace(",", ".")
    elif "." in t:
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})+", t):
            t = t.replace(".", "")
    elif "," in t:
        t = t.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", t) else t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        raise ValueError(f'"{v}" no es un número')


def _read_table(filename, data):
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ApiError('Ese es el formato viejo de Excel (.xls). En Excel usa "Guardar como" → "Libro de Excel (.xlsx)".')
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise ApiError("El servidor no tiene instalado el lector de Excel (openpyxl). Mientras tanto, guarda el archivo como CSV.")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ApiError("No se pudo abrir el archivo de Excel. ¿Está dañado o protegido con contraseña?")
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        return rows
    if name.endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "latin-1"):
            try:
                txt = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(txt[:5000], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        return [r for r in csv.reader(io.StringIO(txt), dialect)]
    raise ApiError("Sube un archivo de Excel (.xlsx) o CSV.")


def _import_plan(conn, filename, data):
    table = _read_table(filename, data)
    header_idx, mapping = None, {}
    for i, row in enumerate(table[:15]):
        found = {}
        for j, h in enumerate(row):
            f = _header_field(h)
            if f and f not in found:
                found[f] = j
        if "name" in found and len(found) >= 2:
            header_idx, mapping = i, found
            break
    if header_idx is None:
        raise ApiError("No se encontró la fila de encabezados. La primera fila debe tener al menos las columnas "
                       "Nombre y Precio (descarga la plantilla para ver el formato).")
    existing = {p["name"].strip().lower(): dict(p) for p in conn.execute("SELECT * FROM products").fetchall()}
    rows, seen = [], set()
    data_rows = table[header_idx + 1:]
    if len(data_rows) > IMPORT_MAX_ROWS:
        raise ApiError(f"El archivo tiene más de {IMPORT_MAX_ROWS} filas. Divídelo en varios archivos.")
    for n, raw in enumerate(data_rows, start=header_idx + 2):
        get = lambda f: raw[mapping[f]] if f in mapping and mapping[f] < len(raw) else None  # noqa: E731
        name = str(get("name") or "").strip()
        values = {f: get(f) for f in ("price", "cost_price", "stock_qty")}
        if not name and all(v in (None, "") for v in values.values()):
            continue  # fila vacía
        entry = {"row": n, "name": name, "errors": [], "changes": []}
        parsed = {}
        for f, label in (("price", "precio"), ("cost_price", "costo"), ("stock_qty", "existencias")):
            try:
                v = _parse_amount(values[f])
            except ValueError as e:
                entry["errors"].append(f"{label}: {e}")
                continue
            if v is not None:
                if v < 0:
                    entry["errors"].append(f"{label} no puede ser negativo")
                parsed[f] = v
        for f in ("category", "unit"):
            v = get(f)
            if v not in (None, ""):
                parsed[f] = str(v).strip()
        key = name.lower()
        if not name:
            entry["errors"].append("falta el nombre")
        elif key in seen:
            entry["errors"].append("el producto está repetido en el archivo")
        seen.add(key)
        current = existing.get(key)
        if current:
            entry["action"], entry["id"] = "actualizar", current["id"]
            if "price" in parsed and parsed["price"] <= 0:
                entry["errors"].append("el precio debe ser mayor que cero")
            labels = {"price": "precio", "cost_price": "costo", "stock_qty": "existencias", "category": "categoría", "unit": "unidad"}
            for f, v in parsed.items():
                old = current.get(f)
                if (isinstance(v, float) and abs((old or 0) - v) > 1e-9) or (isinstance(v, str) and (old or "") != v):
                    entry["changes"].append(f"{labels[f]}: {'—' if old in (None, '') else f'{old:g}'} → {v:g}" if isinstance(v, float)
                                            else f'{labels[f]}: "{old or ""}" → "{v}"')
            if not entry["changes"] and not entry["errors"]:
                entry["action"] = "sin cambios"
        else:
            entry["action"] = "crear"
            if not parsed.get("price") or parsed["price"] <= 0:
                entry["errors"].append("los productos nuevos necesitan un precio mayor que cero")
        entry["values"] = parsed
        rows.append(entry)
    count = lambda a: sum(1 for r in rows if r["action"] == a and not r["errors"])  # noqa: E731
    return {
        "columns": {f: IMPORT_FIELDS[f] for f in mapping},
        "rows": rows,
        "summary": {"crear": count("crear"), "actualizar": count("actualizar"), "sin_cambios": count("sin cambios"),
                    "errores": sum(1 for r in rows if r["errors"])},
    }


@app.post("/api/products/import")
@require_auth(admin=True)
def import_products():
    d = body()
    try:
        data = base64.b64decode(d.get("data_base64") or "", validate=True)
    except Exception:
        raise ApiError("El archivo no llegó bien. Intenta de nuevo.")
    if not data:
        raise ApiError("Elige un archivo.")
    if len(data) > IMPORT_MAX_BYTES:
        raise ApiError("El archivo pesa más de 5 MB.")
    filename = text(d, "filename")
    with get_conn() as conn:
        plan = _import_plan(conn, filename, data)
        if not d.get("apply"):
            return jsonify(plan)
        now = time.time()
        created = updated = 0
        for r in plan["rows"]:
            if r["errors"] or r["action"] == "sin cambios":
                continue
            v = r["values"]
            if r["action"] == "crear":
                pid = str(uuid.uuid4())
                conn.execute(
                    """INSERT INTO products (id, name, price, unit, stock_qty, cost_price, category, created_at)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (pid, r["name"], v["price"], v.get("unit") or "unidad", v.get("stock_qty", 0),
                     v.get("cost_price", 0), v.get("category", ""), now))
                if v.get("stock_qty"):
                    conn.execute("INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                                 (pid, v["stock_qty"], "importación", filename, now))
                created += 1
            else:
                old = conn.execute("SELECT stock_qty FROM products WHERE id=?", (r["id"],)).fetchone()
                sets = {f: v[f] for f in ("price", "cost_price", "stock_qty", "category", "unit") if f in v}
                conn.execute(f"UPDATE products SET {', '.join(f'{f}=?' for f in sets)} WHERE id=?",
                             list(sets.values()) + [r["id"]])
                if "stock_qty" in sets and abs((old["stock_qty"] or 0) - sets["stock_qty"]) > 1e-9:
                    conn.execute("INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                                 (r["id"], sets["stock_qty"] - (old["stock_qty"] or 0), "importación", filename, now))
                log(conn, "producto", r["id"], f'Editó el producto "{r["name"]}" al importar "{filename}": ' + "; ".join(r["changes"]))
                updated += 1
        log(conn, "producto", None, f'Importó productos desde "{filename}": {created} creados, {updated} actualizados'
            + (f', {plan["summary"]["errores"]} filas con errores no se importaron' if plan["summary"]["errores"] else ""))
    return jsonify({"created": created, "updated": updated, "skipped": plan["summary"]["errores"]})


@app.get("/api/products/import-template")
@require_auth(admin=True)
def import_template():
    """Plantilla de Excel con los productos actuales: se puede editar y volver a subir para cambiar precios en bloque."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"
    headers = list(IMPORT_FIELDS.values())
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="16233A")
    with get_conn() as conn:
        products = conn.execute("SELECT * FROM products ORDER BY category COLLATE NOCASE, name COLLATE NOCASE").fetchall()
    for p in products:
        ws.append([p["name"], p["category"] or "", p["price"], p["cost_price"] or 0, p["stock_qty"] or 0, p["unit"] or "unidad"])
    if not products:
        ws.append(["Martillo de uña 16 oz", "Herramientas", 25000, 15000, 10, "unidad"])
        ws.append(["Bombillo LED 9W", "Eléctricos", 8000, 4000, 50, "unidad"])
    for col, width in zip("ABCDEF", (36, 18, 30, 14, 14, 12)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="productos.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- facturas ----------
INVOICE_SELECT = """
SELECT i.*,
  (SELECT COALESCE(SUM(qty * price * (1 - COALESCE(discount, 0) / 100.0)), 0)
     FROM invoice_items WHERE invoice_id = i.id) AS net,
  (SELECT COALESCE(SUM(amount), 0) FROM payments WHERE invoice_id = i.id) AS paid,
  (SELECT COALESCE(SUM(amount), 0) FROM returns WHERE invoice_id = i.id) AS returned
FROM invoices i
"""


def _with_totals(row):
    """Agrega totales y estado de pago a una fila de INVOICE_SELECT.
    total = lo que vale la factura después de devoluciones; original_total = lo facturado."""
    inv = dict(row)
    inv["status"] = inv.get("status") or "activa"
    lines = inv.pop("net")  # suma de las líneas (cantidad × precio − descuento)
    rate = (inv.get("tax_rate") or 0) / 100
    if inv.get("tax_included"):
        inv["original_total"] = money(lines)               # el precio ya trae el IVA
        inv["subtotal"] = money(lines / (1 + rate))         # base sin IVA
    else:
        inv["original_total"] = money(lines * (1 + rate))  # facturas viejas: IVA sumado aparte
        inv["subtotal"] = money(lines)
    inv["tax_amount"] = money(inv["original_total"] - inv["subtotal"])
    inv["returned"] = money(inv["returned"])
    inv["total"] = money(inv["original_total"] - inv["returned"])
    inv["paid"] = money(inv["paid"])
    inv["balance"] = money(inv["total"] - inv["paid"])
    if inv["status"] == "anulada":
        state = "anulada"
    elif inv["returned"] > 0 and inv["total"] <= 0.004:
        state = "devuelta"
    elif not inv.get("payment_tracking"):
        state = "sin_registro"
    elif inv["balance"] <= 0.004:
        state = "pagada"
    elif inv["paid"] > 0:
        state = "abono"
    else:
        state = "pendiente"
    inv["payment_status"] = state
    return inv


def _invoice_rows(conn, where="", params=(), order="ORDER BY i.created_at DESC"):
    return [_with_totals(r) for r in conn.execute(f"{INVOICE_SELECT} {where} {order}", params).fetchall()]


def _load_invoice(conn, iid):
    rows = _invoice_rows(conn, "WHERE i.id=?", (iid,))
    if not rows:
        raise ApiError("Factura no encontrada.", 404)
    inv = rows[0]
    inv["items"] = [dict(i) for i in conn.execute(
        "SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id", (iid,)).fetchall()]
    returned_qty = {r["invoice_item_id"]: r["qty"] for r in conn.execute(
        """SELECT ri.invoice_item_id, SUM(ri.qty) AS qty FROM return_items ri JOIN returns r ON r.id = ri.return_id
           WHERE r.invoice_id=? GROUP BY ri.invoice_item_id""", (iid,)).fetchall()}
    for it in inv["items"]:
        it["returned_qty"] = returned_qty.get(it["id"], 0)
    inv["payments"] = [dict(p) for p in conn.execute(
        "SELECT * FROM payments WHERE invoice_id=? ORDER BY created_at", (iid,)).fetchall()]
    inv["returns"] = []
    for r in conn.execute("SELECT * FROM returns WHERE invoice_id=? ORDER BY created_at", (iid,)).fetchall():
        ret = dict(r)
        ret["items"] = [dict(x) for x in conn.execute("SELECT * FROM return_items WHERE return_id=?", (r["id"],))]
        inv["returns"].append(ret)
    if not is_admin():
        for it in inv["items"]:
            it.pop("unit_cost", None)
        for ret in inv["returns"]:
            ret.pop("cost", None)
    return inv


@app.get("/api/invoices")
@require_auth()
def list_invoices():
    with get_conn() as conn:
        return jsonify(_invoice_rows(conn))


@app.get("/api/invoices/<iid>")
@require_auth()
def get_invoice(iid):
    with get_conn() as conn:
        return jsonify(_load_invoice(conn, iid))


@app.post("/api/invoices")
@require_auth()
def create_invoice():
    d = body()
    client = d.get("client_snapshot") or {}
    items = d.get("items") or []
    tax_rate = num(d.get("tax_rate"), "La tarifa de impuesto", 0, 100, default=0)
    payment = d.get("payment")  # None = app antigua: no se lleva registro de pago
    if payment is not None and not isinstance(payment, dict):
        raise ApiError("La forma de pago no es válida.")
    if not items:
        raise ApiError("La factura necesita al menos un artículo.")
    if not text(client, "name"):
        raise ApiError("La factura necesita un cliente.")
    if payment and payment.get("mode") == "credito" and not client.get("id"):
        raise ApiError("Para vender a crédito elige un cliente registrado (no se le puede fiar a Consumidor final).")

    parsed = []
    for it in items:
        qty = num(it.get("qty"), "La cantidad", 0)
        if qty <= 0:
            raise ApiError("Todas las cantidades deben ser mayores que cero.")
        discount = num(it.get("discount"), "El descuento", 0, 100, default=0)
        price = it.get("price")
        price = None if price is None else num(price, "El precio", 0)
        parsed.append((it.get("product_id"), qty, price, discount))

    invoice_id = str(uuid.uuid4())
    for _attempt in range(5):
        try:
            with get_conn() as conn:
                # Validar producto y stock de TODOS los artículos antes de tocar nada.
                # Si un producto aparece en varias líneas, se suma lo pedido.
                products, wanted = {}, {}
                for pid, qty, _price, _disc in parsed:
                    if pid not in products:
                        product = conn.execute("SELECT * FROM products WHERE id=?", (pid,)).fetchone()
                        if not product:
                            raise ApiError(f"Uno de los productos ya no existe en el catálogo (id {pid}).")
                        products[pid] = product
                    wanted[pid] = wanted.get(pid, 0) + qty
                for pid, qty in wanted.items():
                    p = products[pid]
                    if (p["stock_qty"] or 0) + 1e-9 < qty:
                        raise ApiError(
                            f'No hay suficiente stock de "{p["name"]}": '
                            f'disponible {p["stock_qty"]:g}, solicitado {qty:g}.')

                net = sum(qty * (price if price is not None else products[pid]["price"]) * (1 - disc / 100)
                          for pid, qty, price, disc in parsed)
                total = money(net)  # los precios ya incluyen el IVA
                pay_amount, pay_method = 0.0, None
                if payment is not None:
                    pay_method = parse_method(payment.get("method"))
                    if payment.get("mode") == "credito":
                        pay_amount = money(num(payment.get("amount"), "El abono inicial", 0, default=0))
                        if pay_amount > total:
                            raise ApiError("El abono inicial no puede ser mayor que el total de la factura.")
                    else:
                        pay_amount = total

                row = conn.execute("SELECT invoice_counter FROM company WHERE id=1").fetchone()
                counter = row["invoice_counter"] if row else 1
                number = f"FAC-{counter:04d}"
                now = time.time()
                user = g.user
                client_id = client.get("id")
                if client_id and not conn.execute("SELECT 1 FROM clients WHERE id=?", (client_id,)).fetchone():
                    client_id = None
                conn.execute(
                    """INSERT INTO invoices (id, number, created_at, client_name, client_doc,
                       client_address, client_phone, client_email, tax_rate, notes, warranty, seller,
                       status, user_id, user_name, payment_tracking, client_id, tax_included)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,'activa',?,?,?,?,1)""",
                    (invoice_id, number, now, text(client, "name"), text(client, "doc"),
                     text(client, "address"), text(client, "phone"), text(client, "email"),
                     tax_rate, text(d, "notes"), text(d, "warranty"), text(d, "seller"),
                     user["id"], user["name"], 1 if payment is not None else 0, client_id),
                )
                for pid, qty, price, discount in parsed:
                    product = products[pid]
                    conn.execute(
                        """INSERT INTO invoice_items (invoice_id, product_id, description, qty, price, discount, unit_cost)
                           VALUES (?,?,?,?,?,?,?)""",
                        (invoice_id, pid, product["name"], qty,
                         price if price is not None else product["price"], discount, product["cost_price"] or 0),
                    )
                    conn.execute("UPDATE products SET stock_qty = stock_qty - ? WHERE id=?", (qty, pid))
                    conn.execute(
                        "INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                        (pid, -qty, "venta", number, now),
                    )
                if pay_amount > 0:
                    conn.execute(
                        """INSERT INTO payments (id, invoice_id, amount, method, note, created_at, user_id, user_name)
                           VALUES (?,?,?,?,?,?,?,?)""",
                        (str(uuid.uuid4()), invoice_id, pay_amount, pay_method,
                         "Pago al facturar" if payment.get("mode") != "credito" else "Abono inicial",
                         now, user["id"], user["name"]))
                conn.execute("UPDATE company SET invoice_counter=? WHERE id=1", (counter + 1,))
                mode = "" if payment is None else (" de contado" if payment.get("mode") != "credito"
                                                   else f" a crédito (abono {fmt(pay_amount)})")
                log(conn, "factura", invoice_id,
                    f'Creó la factura {number} para "{text(client, "name")}" por {fmt(total)}{mode}')
            return jsonify({"id": invoice_id, "number": number, "total": total})
        except sqlite3.IntegrityError:
            with get_conn() as conn2:
                row = conn2.execute("SELECT invoice_counter FROM company WHERE id=1").fetchone()
                conn2.execute("UPDATE company SET invoice_counter=? WHERE id=1", ((row["invoice_counter"] if row else 1) + 1,))
            continue
    raise ApiError("No se pudo generar un número de factura único.", 500)


@app.get("/api/invoices/<iid>/pdf")
@require_auth()
def invoice_pdf(iid):
    with get_conn() as conn:
        inv = _load_invoice(conn, iid)
        company = _company(conn)
    buf = io.BytesIO()
    generate_invoice_pdf(inv, company, buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/pdf", download_name=f'{inv["number"]}.pdf')


def _refund_split(conn, table, key, iid, chosen=None):
    """Cómo se devuelve lo pagado: {método: monto}. Por defecto, igual que se pagó (lo neto de cada método);
    si se eligió un método, todo por ese método."""
    nets = {r[0]: money(r[1]) for r in conn.execute(
        f"SELECT method, SUM(amount) FROM {table} WHERE {key}=? GROUP BY method", (iid,)).fetchall()}
    total = money(sum(nets.values()))
    if total <= 0.004:
        return {}
    if chosen:
        return {parse_method(chosen): total}
    split = {m: v for m, v in nets.items() if v > 0.004}
    # si por devoluciones previas algún método quedó negativo, se compensa en el método con más saldo
    extra = money(total - sum(split.values()))
    if extra and split:
        biggest = max(split, key=split.get)
        split[biggest] = money(split[biggest] + extra)
    return {m: v for m, v in split.items() if v > 0.004}


def _split_text(split):
    return " + ".join(f"{fmt(v)} ({m})" for m, v in split.items())


def _void_invoice(iid, reason, refund_method=None):
    reason = (reason or "").strip()
    if not reason:
        raise ApiError("Escribe el motivo de la anulación.")
    with get_conn() as conn:
        inv = conn.execute("SELECT number, status, client_name FROM invoices WHERE id=?", (iid,)).fetchone()
        if not inv:
            raise ApiError("Factura no encontrada.", 404)
        if inv["status"] == "anulada":
            raise ApiError(f'La factura {inv["number"]} ya estaba anulada.')
        now = time.time()
        # se devuelve al inventario lo que no se haya devuelto ya con una devolución parcial
        returned = {r["invoice_item_id"]: r["qty"] for r in conn.execute(
            """SELECT ri.invoice_item_id, SUM(ri.qty) AS qty FROM return_items ri JOIN returns r ON r.id = ri.return_id
               WHERE r.invoice_id=? GROUP BY ri.invoice_item_id""", (iid,)).fetchall()}
        for it in conn.execute("SELECT id, product_id, qty FROM invoice_items WHERE invoice_id=?", (iid,)).fetchall():
            qty = it["qty"] - returned.get(it["id"], 0)
            if it["product_id"] and qty > 0:
                conn.execute("UPDATE products SET stock_qty = stock_qty + ? WHERE id=?", (qty, it["product_id"]))
                conn.execute(
                    "INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                    (it["product_id"], qty, "anulación factura", inv["number"], now))
        # Lo que el cliente ya había pagado se le devuelve: queda como un pago negativo con la fecha de hoy,
        # para que la caja del día descuente ese dinero (y la del día de la venta no cambie).
        split = _refund_split(conn, "payments", "invoice_id", iid, refund_method)
        for method, amount in split.items():
            conn.execute(
                """INSERT INTO payments (id, invoice_id, amount, method, note, created_at, user_id, user_name)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (str(uuid.uuid4()), iid, -amount, method, "Dinero devuelto al anular la factura", now,
                 g.user["id"], g.user["name"]))
        paid = money(sum(split.values()))
        refund_text = f"; se devolvieron {_split_text(split)} al cliente" if split else ""
        conn.execute("UPDATE invoices SET status='anulada', voided_at=?, voided_by=?, void_reason=? WHERE id=?",
                     (now, g.user["name"], reason, iid))
        log(conn, "factura", iid, f'Anuló la factura {inv["number"]} ({inv["client_name"]}): {reason}{refund_text}')
    return jsonify({"ok": True, "refund": paid if paid > 0.004 else 0})


@app.post("/api/invoices/<iid>/void")
@require_auth(admin=True)
def void_invoice(iid):
    d = body()
    return _void_invoice(iid, d.get("reason"), d.get("refund_method"))


@app.delete("/api/invoices/<iid>")
@require_auth(admin=True)
def delete_invoice(iid):
    # Las facturas no se borran (la numeración debe quedar completa): las apps
    # antiguas que llaman a "eliminar" en realidad la anulan.
    return _void_invoice(iid, "Anulada desde una versión anterior de la app")


# ---------- devoluciones parciales ----------
@app.post("/api/invoices/<iid>/returns")
@require_auth(admin=True)
def create_return(iid):
    d = body()
    reason = text(d, "reason")
    if not reason:
        raise ApiError("Escribe el motivo de la devolución.")
    wanted = {}
    for it in d.get("items") or []:
        qty = num(it.get("qty"), "La cantidad devuelta", 0, default=0)
        if qty > 0:
            wanted[int(it.get("item_id"))] = qty
    if not wanted:
        raise ApiError("Indica al menos un artículo y la cantidad que se devuelve.")
    refund_method = parse_method(d.get("refund_method"))
    with get_conn() as conn:
        inv = _load_invoice(conn, iid)
        if inv["status"] == "anulada":
            raise ApiError("La factura está anulada: no se le pueden hacer devoluciones.")
        items = {it["id"]: it for it in inv["items"]}
        lines, amount, net, cost = [], 0.0, 0.0, 0.0
        tax = (inv["tax_rate"] or 0) / 100
        for item_id, qty in wanted.items():
            it = items.get(item_id)
            if not it:
                raise ApiError("Uno de los artículos no pertenece a esta factura.")
            available = it["qty"] - it["returned_qty"]
            if qty > available + 1e-9:
                raise ApiError(f'De "{it["description"]}" solo se pueden devolver {available:g}.')
            line = qty * it["price"] * (1 - (it["discount"] or 0) / 100)
            line_amount = line if inv.get("tax_included") else line * (1 + tax)  # lo que pagó el cliente
            lines.append((it, qty, money(line_amount)))
            net += line_amount / (1 + tax)  # base sin IVA
            amount += line_amount
            cost += qty * (it.get("unit_cost") or 0)
        amount = money(amount)
        # Si el cliente ya pagó más de lo que vale la factura tras la devolución, se le devuelve la diferencia.
        new_total = inv["total"] - amount
        refund = money(max(0.0, inv["paid"] - max(new_total, 0))) if inv["payment_tracking"] else 0.0
        refund = min(refund, amount)
        rid = str(uuid.uuid4())
        now = time.time()
        conn.execute(
            """INSERT INTO returns (id, invoice_id, amount, net, cost, refund, refund_method, reason, created_at, user_id, user_name)
               VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
            (rid, iid, amount, money(net), money(cost), refund, refund_method if refund else None, reason, now,
             g.user["id"], g.user["name"]))
        for it, qty, line_amount in lines:
            conn.execute(
                "INSERT INTO return_items (return_id, invoice_item_id, product_id, description, qty, amount) VALUES (?,?,?,?,?,?)",
                (rid, it["id"], it["product_id"], it["description"], qty, line_amount))
            if it["product_id"]:
                conn.execute("UPDATE products SET stock_qty = stock_qty + ? WHERE id=?", (qty, it["product_id"]))
                conn.execute(
                    "INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                    (it["product_id"], qty, "devolución", inv["number"], now))
        if refund > 0:
            conn.execute(
                """INSERT INTO payments (id, invoice_id, amount, method, note, created_at, user_id, user_name)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (str(uuid.uuid4()), iid, -refund, refund_method, "Dinero devuelto al cliente (devolución)", now,
                 g.user["id"], g.user["name"]))
        detail = ", ".join(f'{qty:g} × {it["description"]}' for it, qty, _a in lines)
        log(conn, "factura", iid, f'Registró una devolución en {inv["number"]} por {fmt(amount)} ({detail}): {reason}'
            + (f"; se devolvieron {fmt(refund)} en {refund_method}" if refund else ""))
        return jsonify(_load_invoice(conn, iid))


# ---------- pagos de facturas ----------
@app.post("/api/invoices/<iid>/payments")
@require_auth()
def add_payment(iid):
    d = body()
    amount = money(num(d.get("amount"), "El monto", 0))
    if amount <= 0:
        raise ApiError("El monto del pago debe ser mayor que cero.")
    method = parse_method(d.get("method"))
    with get_conn() as conn:
        inv = _load_invoice(conn, iid)
        if inv["status"] == "anulada":
            raise ApiError("No se pueden registrar pagos en una factura anulada.")
        limit = inv["balance"] if inv["payment_tracking"] else inv["total"]
        if amount > limit + 0.004:
            raise ApiError(f"El pago supera el saldo pendiente ({fmt(limit)}).")
        conn.execute(
            """INSERT INTO payments (id, invoice_id, amount, method, note, created_at, user_id, user_name)
               VALUES (?,?,?,?,?,?,?,?)""",
            (str(uuid.uuid4()), iid, amount, method, text(d, "note"), time.time(), g.user["id"], g.user["name"]))
        conn.execute("UPDATE invoices SET payment_tracking=1 WHERE id=?", (iid,))
        log(conn, "factura", iid, f'Registró un pago de {fmt(amount)} ({method}) en {inv["number"]}')
        return jsonify(_load_invoice(conn, iid))


@app.delete("/api/payments/<pay_id>")
@require_auth(admin=True)
def delete_payment(pay_id):
    with get_conn() as conn:
        p = conn.execute("""SELECT p.*, i.number FROM payments p JOIN invoices i ON i.id = p.invoice_id
                            WHERE p.id=?""", (pay_id,)).fetchone()
        if p:
            conn.execute("DELETE FROM payments WHERE id=?", (pay_id,))
            log(conn, "factura", p["invoice_id"], f'Quitó un pago de {fmt(p["amount"])} ({p["method"]}) de {p["number"]}')
    return jsonify({"ok": True})


# ---------- compras a proveedores ----------
PURCHASE_SELECT = """
SELECT p.*,
  (SELECT COALESCE(SUM(qty * unit_cost), 0) FROM purchase_items WHERE purchase_id = p.id) AS total,
  (SELECT COALESCE(SUM(amount), 0) FROM supplier_payments WHERE purchase_id = p.id) AS paid
FROM purchases p
"""


def _purchase_totals(row):
    p = dict(row)
    p["total"], p["paid"] = money(p["total"]), money(p["paid"])
    p["balance"] = money(p["total"] - p["paid"]) if p["status"] != "anulada" else 0.0
    if p["status"] == "anulada":
        p["payment_status"] = "anulada"
    elif p["balance"] <= 0.004:
        p["payment_status"] = "pagada"
    elif p["paid"] > 0:
        p["payment_status"] = "abono"
    else:
        p["payment_status"] = "pendiente"
    return p


def _purchase_rows(conn, where="", params=()):
    return [_purchase_totals(r) for r in conn.execute(f"{PURCHASE_SELECT} {where} ORDER BY p.created_at DESC", params)]


def _load_purchase(conn, pid):
    rows = _purchase_rows(conn, "WHERE p.id=?", (pid,))
    if not rows:
        raise ApiError("Compra no encontrada.", 404)
    p = rows[0]
    p["items"] = [dict(r) for r in conn.execute("SELECT * FROM purchase_items WHERE purchase_id=? ORDER BY id", (pid,))]
    p["payments"] = [dict(r) for r in conn.execute(
        "SELECT * FROM supplier_payments WHERE purchase_id=? ORDER BY created_at", (pid,))]
    return p


@app.get("/api/purchases")
@require_auth(admin=True)
def list_purchases():
    with get_conn() as conn:
        if request.args.get("supplier_id"):
            return jsonify(_purchase_rows(conn, "WHERE p.supplier_id=?", (request.args["supplier_id"],)))
        return jsonify(_purchase_rows(conn))


@app.get("/api/purchases/<pid>")
@require_auth(admin=True)
def get_purchase(pid):
    with get_conn() as conn:
        return jsonify(_load_purchase(conn, pid))


@app.post("/api/purchases")
@require_auth(admin=True)
def create_purchase():
    d = body()
    payment = d.get("payment") or {"mode": "contado", "method": "efectivo"}
    if not isinstance(payment, dict):
        raise ApiError("La forma de pago no es válida.")
    lines = []
    for it in d.get("items") or []:
        qty = num(it.get("qty"), "La cantidad", 0)
        if qty <= 0:
            raise ApiError("Todas las cantidades deben ser mayores que cero.")
        lines.append((it.get("product_id"), qty, num(it.get("unit_cost"), "El costo unitario", 0)))
    if not lines:
        raise ApiError("La compra necesita al menos un producto.")
    with get_conn() as conn:
        supplier = conn.execute("SELECT * FROM suppliers WHERE id=?", (d.get("supplier_id"),)).fetchone()
        if not supplier:
            raise ApiError("Elige el proveedor de la compra.")
        products = {}
        for pid, _q, _c in lines:
            p = conn.execute("SELECT * FROM products WHERE id=?", (pid,)).fetchone()
            if not p:
                raise ApiError("Uno de los productos ya no existe en el catálogo.")
            products[pid] = p
        total = money(sum(q * c for _p, q, c in lines))
        method = parse_method(payment.get("method"))
        if payment.get("mode") == "credito":
            pay_amount = money(num(payment.get("amount"), "El abono inicial", 0, default=0))
            if pay_amount > total:
                raise ApiError("El abono inicial no puede ser mayor que el total de la compra.")
        else:
            pay_amount = total
        counter = conn.execute("SELECT purchase_counter FROM company WHERE id=1").fetchone()["purchase_counter"] or 1
        number = f"COM-{counter:04d}"
        pur_id, now = str(uuid.uuid4()), time.time()
        conn.execute(
            """INSERT INTO purchases (id, number, supplier_id, supplier_name, supplier_ref, notes, created_at,
               user_id, user_name, status) VALUES (?,?,?,?,?,?,?,?,?,'activa')""",
            (pur_id, number, supplier["id"], supplier["name"], text(d, "supplier_ref"), text(d, "notes"), now,
             g.user["id"], g.user["name"]))
        for pid, qty, unit_cost in lines:
            # el producto ya pudo cambiar en una línea anterior de esta misma compra
            p = conn.execute("SELECT stock_qty, cost_price FROM products WHERE id=?", (pid,)).fetchone()
            old_stock, old_cost = max(p["stock_qty"] or 0, 0), p["cost_price"] or 0
            # costo promedio ponderado entre lo que había y lo que llega
            new_cost = ((old_stock * old_cost + qty * unit_cost) / (old_stock + qty)) if old_stock > 0 else unit_cost
            conn.execute("INSERT INTO purchase_items (purchase_id, product_id, description, qty, unit_cost) VALUES (?,?,?,?,?)",
                         (pur_id, pid, products[pid]["name"], qty, unit_cost))
            conn.execute("UPDATE products SET stock_qty = stock_qty + ?, cost_price=? WHERE id=?",
                         (qty, round(new_cost, 4), pid))
            conn.execute(
                "INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                (pid, qty, "compra", number, now))
        if pay_amount > 0:
            conn.execute(
                """INSERT INTO supplier_payments (id, purchase_id, amount, method, note, created_at, user_id, user_name)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (str(uuid.uuid4()), pur_id, pay_amount, method,
                 "Pago al comprar" if payment.get("mode") != "credito" else "Abono inicial", now,
                 g.user["id"], g.user["name"]))
        conn.execute("UPDATE company SET purchase_counter=? WHERE id=1", (counter + 1,))
        log(conn, "compra", pur_id, f'Registró la compra {number} a "{supplier["name"]}" por {fmt(total)}'
            + (" de contado" if payment.get("mode") != "credito" else f" a crédito (abono {fmt(pay_amount)})"))
        return jsonify({"id": pur_id, "number": number, "total": total})


@app.post("/api/purchases/<pid>/payments")
@require_auth(admin=True)
def add_supplier_payment(pid):
    d = body()
    amount = money(num(d.get("amount"), "El monto", 0))
    if amount <= 0:
        raise ApiError("El monto del pago debe ser mayor que cero.")
    method = parse_method(d.get("method"))
    with get_conn() as conn:
        p = _load_purchase(conn, pid)
        if p["status"] == "anulada":
            raise ApiError("La compra está anulada.")
        if amount > p["balance"] + 0.004:
            raise ApiError(f"El pago supera lo que se le debe al proveedor ({fmt(p['balance'])}).")
        conn.execute(
            """INSERT INTO supplier_payments (id, purchase_id, amount, method, note, created_at, user_id, user_name)
               VALUES (?,?,?,?,?,?,?,?)""",
            (str(uuid.uuid4()), pid, amount, method, text(d, "note"), time.time(), g.user["id"], g.user["name"]))
        log(conn, "compra", pid, f'Pagó {fmt(amount)} ({method}) al proveedor "{p["supplier_name"]}" por la compra {p["number"]}')
        return jsonify(_load_purchase(conn, pid))


@app.post("/api/purchases/<pid>/void")
@require_auth(admin=True)
def void_purchase(pid):
    d = body()
    reason = text(d, "reason")
    if not reason:
        raise ApiError("Escribe el motivo de la anulación.")
    with get_conn() as conn:
        p = _load_purchase(conn, pid)
        if p["status"] == "anulada":
            raise ApiError("La compra ya estaba anulada.")
        for it in p["items"]:
            prod = conn.execute("SELECT name, stock_qty FROM products WHERE id=?", (it["product_id"],)).fetchone()
            if prod and (prod["stock_qty"] or 0) + 1e-9 < it["qty"]:
                raise ApiError(f'No se puede anular: de "{prod["name"]}" ya no quedan {it["qty"]:g} en inventario '
                               f'(hay {prod["stock_qty"]:g}). Ajusta el stock primero.')
        now = time.time()
        for it in p["items"]:
            conn.execute("UPDATE products SET stock_qty = stock_qty - ? WHERE id=?", (it["qty"], it["product_id"]))
            conn.execute(
                "INSERT INTO stock_movements (product_id, change_qty, reason, reference, created_at) VALUES (?,?,?,?,?)",
                (it["product_id"], -it["qty"], "anulación compra", p["number"], now))
        split = _refund_split(conn, "supplier_payments", "purchase_id", pid, d.get("refund_method"))
        for method, amount in split.items():
            conn.execute(
                """INSERT INTO supplier_payments (id, purchase_id, amount, method, note, created_at, user_id, user_name)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (str(uuid.uuid4()), pid, -amount, method, "Dinero devuelto por el proveedor al anular la compra", now,
                 g.user["id"], g.user["name"]))
        refund_text = f"; el proveedor devolvió {_split_text(split)}" if split else ""
        conn.execute("UPDATE purchases SET status='anulada', voided_at=?, voided_by=?, void_reason=? WHERE id=?",
                     (now, g.user["name"], reason, pid))
        log(conn, "compra", pid, f'Anuló la compra {p["number"]} ({p["supplier_name"]}): {reason}{refund_text}')
    return jsonify({"ok": True})


# ---------- gastos ----------
@app.get("/api/expenses")
@require_auth()
def list_expenses():
    now = time.time()
    start = num(request.args.get("start"), "start", default=now - 31 * 86400)
    end = num(request.args.get("end"), "end", default=now + 86400)
    where, params = "WHERE e.spent_at >= ? AND e.spent_at < ?", [start, end]
    if not is_admin():
        where += " AND e.user_id = ?"
        params.append(g.user["id"])
    with get_conn() as conn:
        rows = conn.execute(f"""SELECT e.*, s.name AS supplier_name FROM expenses e
                                LEFT JOIN suppliers s ON s.id = e.supplier_id {where}
                                ORDER BY e.spent_at DESC""", params).fetchall()
        categories = [r[0] for r in conn.execute(
            "SELECT DISTINCT category FROM expenses WHERE category <> '' ORDER BY category COLLATE NOCASE")]
    return jsonify({"expenses": [dict(r) for r in rows], "categories": categories})


@app.post("/api/expenses")
@require_auth()
def create_expense():
    d = body()
    amount = money(num(d.get("amount"), "El monto", 0))
    if amount <= 0:
        raise ApiError("El monto del gasto debe ser mayor que cero.")
    category = text(d, "category")
    description = text(d, "description")
    if not category and not description:
        raise ApiError("Escribe la categoría o la descripción del gasto.")
    spent_at = num(d.get("spent_at"), "La fecha", default=time.time())
    method = parse_method(d.get("method"))
    eid = str(uuid.uuid4())
    with get_conn() as conn:
        conn.execute(
            """INSERT INTO expenses (id, spent_at, category, description, amount, method, supplier_id, created_at, user_id, user_name)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (eid, spent_at, category, description, amount, method, d.get("supplier_id") or None, time.time(),
             g.user["id"], g.user["name"]))
        log(conn, "gasto", eid, f"Registró un gasto de {fmt(amount)} ({method}): {category or ''} {description}".strip())
    return jsonify({"id": eid})


@app.delete("/api/expenses/<eid>")
@require_auth(admin=True)
def delete_expense(eid):
    with get_conn() as conn:
        e = conn.execute("SELECT * FROM expenses WHERE id=?", (eid,)).fetchone()
        if e:
            conn.execute("DELETE FROM expenses WHERE id=?", (eid,))
            log(conn, "gasto", eid, f'Eliminó un gasto de {fmt(e["amount"])}: {e["category"]} {e["description"]}'.strip())
    return jsonify({"ok": True})


# ---------- cierre de caja ----------
def _cash_scope():
    """Quién: el vendedor solo ve su propia caja; el administrador elige a alguien o 'todos'."""
    if not is_admin():
        return g.user["id"], g.user["name"]
    uid = request.args.get("user_id") or (request.get_json(silent=True) or {}).get("user_id") or ""
    if uid in ("", "all"):
        return None, "Todos"
    with get_conn() as conn:
        u = conn.execute("SELECT name FROM users WHERE id=?", (uid,)).fetchone()
    if not u:
        raise ApiError("Usuario no encontrado.", 404)
    return uid, u["name"]


def _cash_summary(conn, day, tz, scope_uid):
    start, end = day_bounds(day, tz)
    flt, extra = ("AND p.user_id = ?", [scope_uid]) if scope_uid else ("", [])
    sums = lambda sql: {r["method"]: money(r["amount"]) for r in conn.execute(sql, [start, end] + extra)}  # noqa: E731
    received = sums(f"""SELECT p.method, SUM(p.amount) AS amount FROM payments p JOIN invoices i ON i.id = p.invoice_id
                        WHERE p.created_at >= ? AND p.created_at < ? AND p.amount > 0 {flt} GROUP BY p.method""")
    refunded = sums(f"""SELECT p.method, -SUM(p.amount) AS amount FROM payments p
                        WHERE p.created_at >= ? AND p.created_at < ? AND p.amount < 0 {flt} GROUP BY p.method""")
    supplier = sums(f"""SELECT p.method, SUM(p.amount) AS amount FROM supplier_payments p
                        WHERE p.created_at >= ? AND p.created_at < ? AND p.amount > 0 {flt} GROUP BY p.method""")
    supplier_back = sums(f"""SELECT p.method, -SUM(p.amount) AS amount FROM supplier_payments p
                        WHERE p.created_at >= ? AND p.created_at < ? AND p.amount < 0 {flt} GROUP BY p.method""")
    expenses = sums(f"""SELECT p.method, SUM(p.amount) AS amount FROM expenses p
                        WHERE p.spent_at >= ? AND p.spent_at < ? {flt} GROUP BY p.method""")
    expected = {m: money(received.get(m, 0) - refunded.get(m, 0) - supplier.get(m, 0) + supplier_back.get(m, 0)
                         - expenses.get(m, 0)) for m in PAYMENT_METHODS}
    return {"received": received, "refunded": refunded, "supplier_payments": supplier, "supplier_refunds": supplier_back,
            "expenses": expenses, "expected": expected}


@app.get("/api/cash")
@require_auth()
def cash_day():
    day = request.args.get("day") or datetime.date.today().isoformat()
    tz = tz_offset_arg()
    scope_uid, scope_name = _cash_scope()
    with get_conn() as conn:
        summary = _cash_summary(conn, day, tz, scope_uid)
        closing = conn.execute("SELECT * FROM cash_closings WHERE day=? AND COALESCE(scope_user_id,'')=?",
                               (day, scope_uid or "")).fetchone()
    result = {"day": day, "scope_user_id": scope_uid, "scope_name": scope_name, **summary, "closing": None}
    if closing:
        c = dict(closing)
        c["expected"], c["counted"] = _json(c["expected"]), _json(c["counted"])
        result["closing"] = c
    return jsonify(result)


def _json(s):
    return json.loads(s) if s else {}


@app.post("/api/cash/closings")
@require_auth()
def create_closing():
    d = body()
    day = text(d, "day")
    tz = num(d.get("tz"), "tz", -14 * 60, 14 * 60, default=0)
    scope_uid, scope_name = _cash_scope()
    opening = money(num(d.get("opening_cash"), "La base de caja", 0, default=0))
    counted_in = d.get("counted") or {}
    counted = {m: money(num(counted_in.get(m), f"Lo contado en {m}", 0, default=0)) for m in PAYMENT_METHODS}
    with get_conn() as conn:
        summary = _cash_summary(conn, day, tz, scope_uid)
        expected = dict(summary["expected"])
        expected["efectivo"] = money(expected["efectivo"] + opening)
        difference = money(sum(counted.values()) - sum(expected.values()))
        cid = str(uuid.uuid4())
        try:
            conn.execute(
                """INSERT INTO cash_closings (id, day, scope_user_id, scope_name, opening_cash, expected, counted,
                   difference, notes, created_at, user_id, user_name) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                (cid, day, scope_uid, scope_name, opening, json.dumps(expected), json.dumps(counted), difference,
                 text(d, "notes"), time.time(), g.user["id"], g.user["name"]))
        except sqlite3.IntegrityError:
            raise ApiError(f"La caja de {scope_name} del {day} ya se cerró. Un administrador puede borrar ese cierre "
                           "si hay que repetirlo.")
        state = "cuadró" if abs(difference) < 0.005 else (f"sobraron {fmt(difference)}" if difference > 0
                                                          else f"faltaron {fmt(-difference)}")
        log(conn, "caja", cid, f"Cerró la caja de {scope_name} del {day}: {state}")
    return jsonify({"id": cid, "difference": difference, "expected": expected})


@app.get("/api/cash/closings")
@require_auth()
def list_closings():
    where, params = "", []
    if not is_admin():
        where, params = "WHERE scope_user_id = ?", [g.user["id"]]
    with get_conn() as conn:
        rows = [dict(r) for r in conn.execute(
            f"SELECT * FROM cash_closings {where} ORDER BY day DESC, created_at DESC LIMIT 60", params)]
    for r in rows:
        r["expected"], r["counted"] = _json(r["expected"]), _json(r["counted"])
    return jsonify(rows)


@app.delete("/api/cash/closings/<cid>")
@require_auth(admin=True)
def delete_closing(cid):
    with get_conn() as conn:
        c = conn.execute("SELECT * FROM cash_closings WHERE id=?", (cid,)).fetchone()
        if c:
            conn.execute("DELETE FROM cash_closings WHERE id=?", (cid,))
            log(conn, "caja", cid, f'Borró el cierre de caja de {c["scope_name"]} del {c["day"]}')
    return jsonify({"ok": True})


# ---------- registro de actividad (solo administrador) ----------
@app.get("/api/activity")
@require_auth(admin=True)
def activity():
    limit = int(num(request.args.get("limit"), "limit", 1, 500, default=100))
    where, params = [], []
    if request.args.get("before"):
        where.append("created_at < ?")
        params.append(num(request.args.get("before"), "before"))
    q = (request.args.get("q") or "").strip()
    if q:
        where.append("(summary LIKE ? OR user_name LIKE ?)")
        params += [f"%{q}%", f"%{q}%"]
    if request.args.get("entity"):
        where.append("entity = ?")
        params.append(request.args["entity"])
    sql = "SELECT * FROM activity_log" + (" WHERE " + " AND ".join(where) if where else "") + \
          " ORDER BY created_at DESC, id DESC LIMIT ?"
    with get_conn() as conn:
        rows = [dict(r) for r in conn.execute(sql, params + [limit])]
    return jsonify(rows)


# ---------- resumen de ventas (solo administrador) ----------
@app.get("/api/dashboard")
@require_auth(admin=True)
def dashboard():
    now = time.time()
    start = num(request.args.get("start"), "start", default=now - 30 * 86400)
    end = num(request.args.get("end"), "end", default=now)
    tz_offset = tz_offset_arg()

    def local_day(ts):
        return datetime.datetime.fromtimestamp(ts - tz_offset * 60, datetime.timezone.utc).date()

    active = "COALESCE(i.status,'activa') <> 'anulada'"
    with get_conn() as conn:
        invoices = _invoice_rows(conn, f"WHERE i.created_at >= ? AND i.created_at < ? AND {active}", (start, end))
        items = conn.execute(
            f"""SELECT it.*, COALESCE(it.unit_cost, p.cost_price, 0) AS cost, i.tax_rate AS tax_rate,
                      i.tax_included AS tax_included FROM invoice_items it
               JOIN invoices i ON i.id = it.invoice_id LEFT JOIN products p ON p.id = it.product_id
               WHERE i.created_at >= ? AND i.created_at < ? AND {active}""", (start, end)).fetchall()
        rets = conn.execute(
            f"""SELECT r.* FROM returns r JOIN invoices i ON i.id = r.invoice_id
                WHERE r.created_at >= ? AND r.created_at < ? AND {active}""", (start, end)).fetchall()
        ret_items = conn.execute(
            f"""SELECT ri.product_id, ri.description, ri.qty, ri.amount, i.tax_rate FROM return_items ri
                JOIN returns r ON r.id = ri.return_id JOIN invoices i ON i.id = r.invoice_id
                WHERE r.created_at >= ? AND r.created_at < ? AND {active}""", (start, end)).fetchall()
        payments = conn.execute(
            """SELECT p.method, SUM(p.amount) AS amount FROM payments p
               WHERE p.created_at >= ? AND p.created_at < ?
               GROUP BY p.method HAVING ABS(SUM(p.amount)) > 0.004 ORDER BY amount DESC""", (start, end)).fetchall()
        receivables = [r for r in _invoice_rows(conn, f"WHERE {active} AND i.payment_tracking = 1", order="ORDER BY i.created_at")
                       if r["balance"] > 0.004]
        expenses = conn.execute(
            """SELECT COALESCE(NULLIF(category,''),'Sin categoría') AS category, SUM(amount) AS amount FROM expenses
               WHERE spent_at >= ? AND spent_at < ? GROUP BY 1 ORDER BY amount DESC""", (start, end)).fetchall()
        purchases = _purchase_rows(conn, "WHERE p.created_at >= ? AND p.created_at < ? AND p.status <> 'anulada'", (start, end))
        payables = [p for p in _purchase_rows(conn, "WHERE p.status <> 'anulada'") if p["balance"] > 0.004]

    def base(i, with_discount=True):
        """Valor de la línea sin IVA (si el precio traía el IVA, se le quita)."""
        v = i["qty"] * i["price"] * ((1 - (i["discount"] or 0) / 100) if with_discount else 1)
        return v / (1 + (i["tax_rate"] or 0) / 100) if i["tax_included"] else v

    gross = sum(base(i, with_discount=False) for i in items)
    net = sum(base(i) for i in items)
    cost = sum(i["qty"] * (i["cost"] or 0) for i in items)
    total = sum(inv["original_total"] for inv in invoices)
    ret_total = sum(r["amount"] for r in rets)
    ret_net = sum(r["net"] for r in rets)
    ret_cost = sum(r["cost"] or 0 for r in rets)
    expense_total = sum(e["amount"] for e in expenses)
    sales_net = net - ret_net
    profit = sales_net - (cost - ret_cost)

    products = {}
    for i in items:
        key = i["product_id"] or i["description"]
        p = products.setdefault(key, {"name": i["description"], "qty": 0, "net": 0.0})
        p["qty"] += i["qty"]
        p["net"] += base(i)
    for r in ret_items:
        key = r["product_id"] or r["description"]
        if key in products:
            products[key]["qty"] -= r["qty"]
            products[key]["net"] -= r["amount"] / (1 + (r["tax_rate"] or 0) / 100)

    sellers = {}
    for inv in invoices:
        name = (inv.get("seller") or "").strip() or inv.get("user_name") or "Sin vendedor"
        s = sellers.setdefault(name, {"name": name, "count": 0, "total": 0.0})
        s["count"] += 1
        s["total"] += inv["original_total"]

    # serie por día (o por mes si el período es largo), con ceros en los días sin ventas
    first, last = local_day(start), local_day(max(start, end - 1))
    monthly = (last - first).days > 62
    buckets = {}
    day = first
    while day <= last:
        key = day.strftime("%Y-%m") if monthly else day.isoformat()
        buckets.setdefault(key, {"date": key, "total": 0.0, "count": 0})
        day += datetime.timedelta(days=1)
    for inv in invoices:
        d = local_day(inv["created_at"])
        key = d.strftime("%Y-%m") if monthly else d.isoformat()
        if key in buckets:
            buckets[key]["total"] += inv["original_total"]
            buckets[key]["count"] += 1

    return jsonify({
        "summary": {
            "count": len(invoices), "gross": money(gross), "discount": money(gross - net),
            "net": money(sales_net), "tax": money(total - net - (ret_total - ret_net)),
            "total": money(total - ret_total), "returns": money(ret_total), "returns_count": len(rets),
            "cost": money(cost - ret_cost), "profit": money(profit),
            "expenses": money(expense_total), "net_profit": money(profit - expense_total),
            "avg_ticket": money(total / len(invoices)) if invoices else 0,
            "items_without_cost": sum(1 for i in items if not i["cost"]),
        },
        "series": {"unit": "month" if monthly else "day",
                   "points": [dict(b, total=money(b["total"])) for b in buckets.values()]},
        "collected": {"total": money(sum(p["amount"] for p in payments)),
                      "by_method": [{"method": p["method"], "amount": money(p["amount"])} for p in payments]},
        "top_products": sorted(({**p, "net": money(p["net"])} for p in products.values() if p["qty"] > 1e-9),
                               key=lambda p: -p["net"])[:10],
        "by_seller": sorted(({**s, "total": money(s["total"])} for s in sellers.values()),
                            key=lambda s: -s["total"]),
        "expenses_by_category": [{"category": e["category"], "amount": money(e["amount"])} for e in expenses],
        "purchases": {"total": money(sum(p["total"] for p in purchases)), "count": len(purchases)},
        "payables": {"total": money(sum(p["balance"] for p in payables)), "count": len(payables)},
        "receivables": {
            "total": money(sum(r["balance"] for r in receivables)), "count": len(receivables),
            "invoices": [{k: r[k] for k in ("id", "number", "client_name", "client_id", "created_at", "total", "balance")}
                         for r in receivables[:50]],
        },
    })


# ---------- copias de seguridad ----------
def _write_backup(dest: Path):
    """Copia consistente de la base (aunque haya gente usándola) usando la API de backup de SQLite."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".db", dir=dest.parent)
    os.close(fd)
    try:
        src = sqlite3.connect(DB_PATH, timeout=10)
        dst = sqlite3.connect(tmp)
        with dst:
            src.backup(dst)
        dst.close()
        src.close()
        os.replace(tmp, dest)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


_last_backup_check = 0.0


@app.before_request
def auto_backup():
    """Una vez al día (aprovechando cualquier petición) guarda una copia en backups/."""
    global _last_backup_check
    now = time.time()
    if now - _last_backup_check < 3600:
        return
    _last_backup_check = now
    try:
        autos = sorted(BACKUP_DIR.glob("auto-*.db")) if BACKUP_DIR.exists() else []
        if autos and now - autos[-1].stat().st_mtime < AUTO_BACKUP_EVERY:
            return
        _write_backup(BACKUP_DIR / f"auto-{datetime.datetime.now():%Y-%m-%d_%H%M}.db")
        for old in sorted(BACKUP_DIR.glob("auto-*.db"))[:-AUTO_BACKUP_KEEP]:
            old.unlink()
    except Exception:
        app.logger.exception("No se pudo crear la copia de seguridad automática")


@app.get("/api/backup")
@require_auth(admin=True)
def download_backup():
    fd, tmp = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        src = sqlite3.connect(DB_PATH, timeout=10)
        dst = sqlite3.connect(tmp)
        with dst:
            src.backup(dst)
        dst.close()
        src.close()
        with open(tmp, "rb") as f:
            data = io.BytesIO(f.read())
    finally:
        os.remove(tmp)
    with get_conn() as conn:
        log(conn, "empresa", "1", "Descargó una copia de seguridad")
    return send_file(data, mimetype="application/octet-stream", as_attachment=True,
                     download_name=f"facturador-copia-{datetime.datetime.now():%Y-%m-%d_%H%M}.db")


@app.get("/api/backups")
@require_auth(admin=True)
def list_backups():
    files = sorted(BACKUP_DIR.glob("auto-*.db"), reverse=True) if BACKUP_DIR.exists() else []
    return jsonify([{"name": f.name, "size": f.stat().st_size, "created_at": f.stat().st_mtime} for f in files])


@app.get("/api/backups/<name>")
@require_auth(admin=True)
def get_backup(name):
    if not (name.startswith("auto-") and name.endswith(".db")) or "/" in name or "\\" in name or ".." in name:
        raise ApiError("Copia no encontrada.", 404)
    path = BACKUP_DIR / name
    if not path.is_file():
        raise ApiError("Copia no encontrada.", 404)
    return send_file(path, mimetype="application/octet-stream", as_attachment=True, download_name=name)


# Al final del módulo: init_db() usa funciones definidas más arriba (p. ej. _refund_split).
init_db()


if __name__ == "__main__":
    app.run(debug=True, port=5000)
